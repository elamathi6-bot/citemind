import os, re, json, time, hashlib, secrets, sqlite3, tempfile
import chromadb
import streamlit as st
from dotenv import load_dotenv
from google import genai
from google.genai import types
from pypdf import PdfReader
from pptx import Presentation
from sentence_transformers import SentenceTransformer

load_dotenv()
MODEL = "gemini-3.8-flash"
FALLBACK_MODELS = ["gemini-3.1-flash-lite-preview"]
MAX_DISTANCE = 1.3
TOP_K = 10
REFUSAL = "This isn't covered in your course material."
VIDEO_EXT = {"mp4", "mov", "mkv", "webm", "mp3", "wav", "m4a"}

st.set_page_config(page_title="Study Companion", page_icon="📚", layout="centered")

# ------------------------------------------------------------ look and feel

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Poppins:wght@400;500;600;700&display=swap');
html, body, .stApp, .stMarkdown, p, label, input, textarea, button, h1, h2, h3, h4 { font-family: 'Poppins', sans-serif; }
[data-testid="stIconMaterial"], .material-icons, .material-symbols-rounded { font-family: 'Material Symbols Rounded' !important; }
input, textarea { background:#F3F1FF !important; border-radius:10px !important; color:#1F1B3A !important; }
.stApp { background: linear-gradient(180deg, #F1EEFF 0%, #FFFFFF 40%); }
.hero { display:flex; align-items:center; gap:20px; padding:22px 26px; margin-bottom:18px;
  border-radius:22px; color:white;
  background: linear-gradient(135deg, #6C5CE7 0%, #A29BFE 55%, #74B9FF 100%);
  box-shadow: 0 10px 30px rgba(108,92,231,.35); }
.hero h1 { margin:0; font-size:1.9rem; font-weight:700; color:white; }
.hero p { margin:4px 0 0; opacity:.92; font-size:.95rem; }
.stButton > button, .stFormSubmitButton > button, .stDownloadButton > button {
  border-radius:12px; border:0; font-weight:600; color:white;
  background: linear-gradient(135deg, #6C5CE7, #8E7CFF); }
.stButton > button:hover, .stFormSubmitButton > button:hover { filter:brightness(1.1); color:white; }
[data-baseweb="tab-list"] { gap:6px; }
[data-baseweb="tab"] { border-radius:12px 12px 0 0; font-weight:600; }
[data-testid="stExpander"] { border-radius:14px; background:white; box-shadow:0 2px 10px rgba(0,0,0,.05); }
[data-testid="stForm"] { border-radius:16px; background:white; box-shadow:0 2px 12px rgba(108,92,231,.10); }
#MainMenu, footer { visibility:hidden; }
</style>
""", unsafe_allow_html=True)

HERO_SVG = """<svg width="64" height="64" viewBox="0 0 64 64" fill="none">
<rect x="8" y="12" width="22" height="40" rx="4" fill="white" opacity=".95"/>
<rect x="34" y="12" width="22" height="40" rx="4" fill="white" opacity=".75"/>
<path d="M14 22h10M14 29h10M14 36h6M40 22h10M40 29h10M40 36h6" stroke="#6C5CE7" stroke-width="3" stroke-linecap="round"/>
<circle cx="52" cy="14" r="7" fill="#FFEAA7"/></svg>"""


def hero(sub):
    st.markdown(f'<div class="hero">{HERO_SVG}<div><h1>Study Companion</h1><p>{sub}</p></div></div>',
                unsafe_allow_html=True)


# ------------------------------------------------------------ accounts and saved items

def db():
    con = sqlite3.connect("study.db")
    con.execute("CREATE TABLE IF NOT EXISTS users(username TEXT PRIMARY KEY, salt TEXT, pw TEXT)")
    con.execute("""CREATE TABLE IF NOT EXISTS saved(id INTEGER PRIMARY KEY AUTOINCREMENT,
                   username TEXT, kind TEXT, title TEXT, content TEXT,
                   created TEXT DEFAULT CURRENT_TIMESTAMP)""")
    return con


def hash_pw(pw, salt):
    return hashlib.pbkdf2_hmac("sha256", pw.encode(), bytes.fromhex(salt), 200_000).hex()


def signup(user, pw):
    user = user.strip().lower()
    if not re.fullmatch(r"[a-z0-9_]{3,20}", user):
        return "Username: 3 to 20 letters, numbers or underscores."
    if len(pw) < 6:
        return "Password must be at least 6 characters."
    salt = secrets.token_hex(16)
    try:
        with db() as con:
            con.execute("INSERT INTO users VALUES (?,?,?)", (user, salt, hash_pw(pw, salt)))
    except sqlite3.IntegrityError:
        return "That username is taken."
    return None


def login(user, pw):
    user = user.strip().lower()
    row = db().execute("SELECT salt, pw FROM users WHERE username=?", (user,)).fetchone()
    return bool(row) and secrets.compare_digest(row[1], hash_pw(pw, row[0]))


def save_item(kind, title, content):
    with db() as con:
        con.execute("INSERT INTO saved(username, kind, title, content) VALUES (?,?,?,?)",
                    (st.session_state.user, kind, title, content))
    st.toast("Saved to your library ✅")


# ------------------------------------------------------------ AI helpers

@st.cache_resource
def load_resources():
    llm = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))
    embedder = SentenceTransformer("all-MiniLM-L6-v2")
    client = chromadb.PersistentClient(path="chroma_db")
    return llm, embedder, client


llm, embedder, chroma_client = load_resources()


def label(m):
    src = m.get("source", "Course")
    if m.get("type") == "video":
        mins, secs = divmod(int(m["timestamp"]), 60)
        return f"{src} {mins}:{secs:02d}"
    return f"{src} p.{m.get('page', '?')}"


def retrieve(question, k=TOP_K):
    if col.count() == 0:
        return []
    q = embedder.encode([question]).tolist()
    res = col.query(query_embeddings=q, n_results=min(k, col.count()))
    return list(zip(res["documents"][0], res["metadatas"][0], res["distances"][0]))


def build_context(hits):
    return "\n\n".join(f"[{i+1}] ({label(m)})\n{doc}" for i, (doc, m, _) in enumerate(hits))


def generate(contents, json_mode=False):
    config = types.GenerateContentConfig(response_mime_type="application/json") if json_mode else None
    for model_name in [MODEL] + FALLBACK_MODELS:
        for attempt in range(3):
            try:
                resp = llm.models.generate_content(model=model_name, contents=contents, config=config)
                return resp.text, model_name
            except Exception as e:
                msg = str(e)
                if "503" in msg or "429" in msg:
                    time.sleep(5 * (attempt + 1))
                elif "404" in msg:
                    break
                else:
                    raise
    return None, None


def split_text(text, size=1200):
    text = " ".join(text.split())
    return [text[i:i + size] for i in range(0, len(text), size)] if text else []


def store(chunks, source):
    if not chunks:
        return 0
    texts = [c[0] for c in chunks]
    metas = [{**m, "source": source} for _, m in chunks]
    ids = [hashlib.md5(f"{source}-{i}".encode()).hexdigest() for i in range(len(chunks))]
    col.upsert(ids=ids, documents=texts, metadatas=metas, embeddings=embedder.encode(texts).tolist())
    return len(chunks)


def ingest_pdf(path, name):
    chunks, empty = [], 0
    for pno, page in enumerate(PdfReader(path).pages, 1):
        pieces = split_text(page.extract_text() or "")
        empty += not pieces
        chunks += [(p, {"type": "pdf", "page": pno}) for p in pieces]
    return store(chunks, name), empty


def ingest_pptx(path, name):
    chunks, empty = [], 0
    for sno, slide in enumerate(Presentation(path).slides, 1):
        parts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        if slide.has_notes_slide:
            parts.append(slide.notes_slide.notes_text_frame.text)
        pieces = split_text(" ".join(parts))
        empty += not pieces
        chunks += [(p, {"type": "slide", "page": sno}) for p in pieces]
    return store(chunks, name), empty


def ingest_video(path, name):
    up = llm.files.upload(file=path)
    while up.state.name == "PROCESSING":
        time.sleep(3)
        up = llm.files.get(name=up.name)
    if up.state.name != "ACTIVE":
        raise RuntimeError(f"Gemini could not process this file (state {up.state.name}).")
    prompt = """Transcribe the speech in this recording.
Return ONLY a JSON list of segments of roughly 30 to 60 seconds each:
[{"start": <start time in whole seconds>, "text": "<what was said>"}]
Use only what is actually said. Do not summarize or add anything."""
    try:
        text, _ = generate([up, prompt], json_mode=True)
    finally:
        try:
            llm.files.delete(name=up.name)
        except Exception:
            pass
    if text is None:
        raise RuntimeError("All models are busy. Try again in a few minutes.")
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    chunks = []
    for seg in json.loads(text):
        try:
            body = " ".join(str(seg["text"]).split())
            if body:
                chunks.append((body, {"type": "video", "timestamp": int(float(seg["start"]))}))
        except (KeyError, TypeError, ValueError):
            continue
    return store(chunks, name), 0


def list_sources():
    counts = {}
    for m in col.get(include=["metadatas"])["metadatas"]:
        s = m.get("source", "Course")
        counts[s] = counts.get(s, 0) + 1
    return counts


def grounded(task, hits):
    text, model = generate(f"""You are a study tutor. {task}
Rules:
- Use ONLY the sources below. Do not use outside knowledge.
- Cite sources inline like [1] or [2] after each claim.
- Sources may use different words than the question (for example "put" vs "insert"). Treat them as the same idea if the text or code clearly does that job.
- If a source contains code that answers the question, explain it step by step.
- Ignore loosely related sources. Keep it short and clear.
- If the sources do not contain the answer, reply exactly: "{REFUSAL}"

SOURCES:
{build_context(hits)}""")
    if text is None:
        return None, [], None
    text = text.strip()
    if "isn't covered in your course material" in text.lower():
        return REFUSAL, [], model
    cited = sorted({int(n) for g in re.findall(r"\[([\d,\s]+)\]", text) for n in re.findall(r"\d+", g)})
    return text, [(n, *hits[n - 1]) for n in cited if 1 <= n <= len(hits)], model


def ask(question):
    hits = retrieve(question)
    if not hits or hits[0][2] > MAX_DISTANCE:
        return REFUSAL, [], None
    return grounded(f"Answer the student's question: {question}", hits)


def summarize(topic):
    hits = retrieve(topic, k=12)
    if not hits or hits[0][2] > MAX_DISTANCE:
        return REFUSAL, [], None
    return grounded(f'Write a study summary of "{topic}" as 5 to 8 short bullet points.', hits)


def parse_quiz(text, num_sources):
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        return None
    if isinstance(data, dict):
        data = data.get("questions", [])
    good = []
    for q in data:
        try:
            ans = str(q["answer"]).strip().upper()[:1]
            if len(q["options"]) == 4 and ans in "ABCD" and q["question"] and q["explanation"]:
                src = int(q.get("source", 0))
                good.append({"question": q["question"], "options": q["options"], "answer": ans,
                             "explanation": q["explanation"],
                             "source": src if 1 <= src <= num_sources else None})
        except (KeyError, TypeError, ValueError):
            continue
    return good or None


def make_quiz(topic, n, difficulty):
    hits = retrieve(topic)
    if not hits or hits[0][2] > MAX_DISTANCE:
        return None, hits, "not_covered"
    prompt = f"""You are a study tutor. Write {n} multiple-choice questions about "{topic}"
using ONLY the sources below. Do not use outside knowledge. Difficulty: {difficulty}.
Rules:
- Each question has exactly 4 options and exactly one correct answer.
- Wrong options must be plausible but clearly wrong according to the sources.
- Mix question types: facts, how a process or code works, and why something is done.
- Do not ask two questions that test the same fact.
- "source" is the number [n] of the source that supports the correct answer.
- Return ONLY a JSON list, no other text, in this format:
[{{"question": "...", "options": ["...", "...", "...", "..."], "answer": "B", "explanation": "one or two sentences", "source": 3}}]

SOURCES:
{build_context(hits)}"""
    for _ in range(2):
        text, _m = generate(prompt, json_mode=True)
        if text is None:
            return None, hits, "busy"
        questions = parse_quiz(text, len(hits))
        if questions:
            return questions, hits, "ok"
    return None, hits, "malformed"


def run_and_store(key, title, fn, *args):
    st.session_state[key] = (title, *fn(*args))


def render_result(key, kind):
    r = st.session_state.get(key)
    if not r:
        return
    title, text, sources, model = r
    if text is None:
        st.error("All models are busy right now. Try again in a few minutes.")
        return
    if text == REFUSAL:
        st.warning(REFUSAL)
        return
    st.markdown(text)
    st.subheader("Sources")
    for n, doc, m, dist in sources:
        with st.expander(f"[{n}] {label(m)}  (distance {dist:.2f})"):
            st.text(doc)
    if model:
        st.caption(f"Answered by {model}")
    content = text + "\n\nSources:\n" + "\n".join(f"[{n}] {label(m)}" for n, _, m, _ in sources)
    st.button("💾 Save this", key=f"save_{key}", on_click=save_item, args=(kind, title, content))


# ------------------------------------------------------------ login gate

if "user" not in st.session_state:
    hero("Sign in to keep your material and saved notes private.")
    t_in, t_up = st.tabs(["Log in", "Sign up"])
    with t_in:
        with st.form("login_form"):
            u = st.text_input("Username")
            p = st.text_input("Password", type="password")
            if st.form_submit_button("Log in"):
                if login(u, p):
                    st.session_state.user = u.strip().lower()
                    st.rerun()
                else:
                    st.error("Wrong username or password.")
    with t_up:
        with st.form("signup_form"):
            u2 = st.text_input("Choose a username")
            p2 = st.text_input("Choose a password", type="password")
            if st.form_submit_button("Create account"):
                err = signup(u2, p2)
                if err:
                    st.error(err)
                else:
                    st.session_state.user = u2.strip().lower()
                    st.rerun()
    st.stop()

USER = st.session_state.user
col = chroma_client.get_or_create_collection("u_" + hashlib.md5(USER.encode()).hexdigest()[:16])

with st.sidebar:
    st.markdown(f"### 👋 {USER}")
    st.caption("Your uploads and saved items are private to this account.")
    if st.button("Log out"):
        st.session_state.clear()
        st.rerun()

hero("Upload your course material. Answers come only from it, with citations.")

tab_up, tab_ask, tab_sum, tab_quiz, tab_saved = st.tabs(
    ["📤 Upload", "💬 Ask", "📝 Summary", "🎯 Quiz", "⭐ Saved"]
)

# ---- Upload
with tab_up:
    files = st.file_uploader("PDF, PowerPoint (.pptx), or video/audio",
                             type=["pdf", "pptx"] + sorted(VIDEO_EXT), accept_multiple_files=True)
    if st.button("Add to my course", disabled=not files):
        for f in files:
            ext = f.name.rsplit(".", 1)[-1].lower()
            with st.status(f"Processing {f.name}...", expanded=True) as status:
                tmp = tempfile.NamedTemporaryFile(delete=False, suffix="." + ext)
                try:
                    tmp.write(f.getbuffer())
                    tmp.close()
                    if ext == "pdf":
                        n, empty = ingest_pdf(tmp.name, f.name)
                    elif ext == "pptx":
                        n, empty = ingest_pptx(tmp.name, f.name)
                    else:
                        st.write("Transcribing with Gemini. Long videos can take a few minutes.")
                        n, empty = ingest_video(tmp.name, f.name)
                    msg = f"Stored {n} chunks."
                    if empty:
                        msg += f" {empty} pages had no readable text (scanned images?)."
                    status.update(label=f"{f.name}: {msg}", state="complete")
                except Exception as e:
                    status.update(label=f"{f.name} failed: {e}", state="error")
                finally:
                    try:
                        os.unlink(tmp.name)
                    except OSError:
                        pass
    st.divider()
    st.subheader("Your course material")
    sources = list_sources()
    if not sources:
        st.info("Nothing uploaded yet.")
    for s, count in sources.items():
        c1, c2 = st.columns([4, 1])
        c1.write(f"**{s}**: {count} chunks")
        if c2.button("Delete", key=f"del_{s}"):
            col.delete(where={"source": s})
            st.rerun()
    if sources and st.button("Clear everything"):
        chroma_client.delete_collection(col.name)
        st.rerun()

# ---- Ask
with tab_ask:
    with st.form("ask_form"):
        question = st.text_input("Your question", placeholder="What are the rules of a binary search tree?")
        go = st.form_submit_button("Ask")
    if go and question.strip():
        if col.count() == 0:
            st.warning("Upload some material first.")
        else:
            with st.spinner("Searching your material..."):
                run_and_store("res_ask", question.strip(), ask, question.strip())
    render_result("res_ask", "Answer")

# ---- Summary
with tab_sum:
    with st.form("sum_form"):
        sum_topic = st.text_input("Topic to summarize", placeholder="binary search tree")
        go_sum = st.form_submit_button("Summarize")
    if go_sum and sum_topic.strip():
        if col.count() == 0:
            st.warning("Upload some material first.")
        else:
            with st.spinner("Writing summary..."):
                run_and_store("res_sum", sum_topic.strip(), summarize, sum_topic.strip())
    render_result("res_sum", "Summary")

# ---- Quiz
with tab_quiz:
    with st.form("topic_form"):
        topic = st.text_input("Quiz topic", placeholder="binary search tree")
        c1, c2 = st.columns(2)
        num_q = c1.slider("Number of questions", 3, 10, 5)
        diff = c2.selectbox("Difficulty", ["easy", "medium", "hard"], index=1)
        make = st.form_submit_button("Generate quiz")
    if make and topic.strip():
        if col.count() == 0:
            st.warning("Upload some material first.")
        else:
            with st.spinner("Writing questions..."):
                questions, hits, status = make_quiz(topic.strip(), num_q, diff)
            if status == "ok":
                st.session_state.qid = st.session_state.get("qid", 0) + 1
                st.session_state.quiz = {"questions": questions, "hits": hits, "done": False,
                                         "topic": topic.strip()}
            else:
                st.session_state.pop("quiz", None)
                st.warning({"not_covered": "This topic isn't covered in your course material.",
                            "busy": "All models are busy right now. Try again in a few minutes.",
                            "malformed": "Could not generate a valid quiz. Try again or pick another topic."}[status])

    quiz = st.session_state.get("quiz")
    if quiz:
        questions, hits, qid = quiz["questions"], quiz["hits"], st.session_state.qid
        with st.form(f"quiz_form_{qid}"):
            for i, q in enumerate(questions):
                st.markdown(f"**Q{i+1}. {q['question']}**")
                st.radio("Choose one", [f"{L}) {o}" for L, o in zip("ABCD", q["options"])],
                         index=None, key=f"q{qid}_{i}", label_visibility="collapsed")
            if st.form_submit_button("Submit answers"):
                quiz["done"] = True
        if quiz["done"]:
            score, lines = 0, []
            for i, q in enumerate(questions):
                picked = st.session_state.get(f"q{qid}_{i}")
                ok = bool(picked) and picked[0] == q["answer"]
                score += ok
                src = f" (Source: {label(hits[q['source'] - 1][1])})" if q["source"] else ""
                lines.append(f"Q{i+1}. {q['question']}\nCorrect answer: {q['answer']}. "
                             f"You chose: {picked[0] if picked else 'none'}.\n{q['explanation']}{src}\n")
                with st.expander(f"{'✅' if ok else '❌'} Q{i+1}: correct answer {q['answer']}"):
                    st.write(q["explanation"])
                    if src:
                        st.caption(src.strip(" ()"))
            st.success(f"Score: {score}/{len(questions)}")
            st.button("💾 Save quiz and results", key=f"save_quiz_{qid}", on_click=save_item,
                      args=("Quiz", f"Quiz: {quiz['topic']} ({score}/{len(questions)})",
                            f"Score: {score}/{len(questions)}\n\n" + "\n".join(lines)))

# ---- Saved
with tab_saved:
    rows = db().execute("SELECT id, kind, title, content, created FROM saved WHERE username=? "
                        "ORDER BY id DESC", (USER,)).fetchall()
    if not rows:
        st.info("Nothing saved yet. Use the 💾 Save button on an answer, summary or quiz.")
    for rid, kind, title, content, created in rows:
        with st.expander(f"{kind}: {title[:70]}  ·  {created[:16]}"):
            if kind == "Quiz":
                st.text(content)
            else:
                st.markdown(content)
            d1, d2 = st.columns(2)
            d1.download_button("Download", content, file_name=f"{kind}_{rid}.txt", key=f"dl_{rid}")
            if d2.button("Delete", key=f"rm_{rid}"):
                with db() as con:
                    con.execute("DELETE FROM saved WHERE id=? AND username=?", (rid, USER))
                st.rerun()